import os
import sys
sys.path.insert(0, 'src')

try:
    from docx import Document
except ImportError:
    print("请先安装 python-docx: pip install python-docx")
    sys.exit(1)


def create_contract_1(output_path):
    doc = Document()
    doc.add_heading('房屋买卖合同', 0)

    doc.add_heading('第一条 定义', level=1)
    doc.add_paragraph('1.1 甲方：指出售房屋的一方，即张三。')
    doc.add_paragraph('1.2 乙方：指购买房屋的一方，即李四。')
    doc.add_paragraph('1.3 房屋：指位于北京市朝阳区某某路123号的房屋。')

    doc.add_heading('第二条 合同价款', level=1)
    doc.add_paragraph('2.1 本房屋的总价款为人民币500万元整。')
    doc.add_paragraph('2.2 乙方应在签订合同之日支付定金50万元。')
    doc.add_paragraph('2.3 剩余款项应在办理过户手续前一次性付清。')

    doc.add_heading('第三条 违约责任', level=1)
    doc.add_paragraph('3.1 任何一方违反本合同约定，应承担违约责任。')
    doc.add_paragraph('3.2 违约方应向守约方支付违约金，违约金金额为合同总金额的10%。')
    doc.add_paragraph('3.3 若违约金不足以弥补守约方损失的，违约方还应赔偿差额部分。')

    doc.add_heading('第四条 争议解决', level=1)
    doc.add_paragraph('4.1 因本合同引起的或与本合同有关的任何争议，双方应友好协商解决。')
    doc.add_paragraph('4.2 协商不成的，任何一方均有权向房屋所在地有管辖权的人民法院提起诉讼。')

    doc.save(output_path)
    print(f"已创建: {output_path}")


def create_contract_2(output_path):
    doc = Document()
    doc.add_heading('服务合同', 0)

    doc.add_heading('第一条 服务内容', level=1)
    doc.add_paragraph('1.1 乙方同意按照本合同约定，向甲方提供技术咨询服务。')
    doc.add_paragraph('1.2 服务期限为自合同生效之日起一年。')

    doc.add_heading('第二条 付款方式', level=1)
    doc.add_paragraph('2.1 本项目总费用为人民币50万元整。')
    doc.add_paragraph('2.2 甲方应于每月5日前支付当月服务费。')
    doc.add_paragraph('2.3 逾期支付的，每日按应付金额的0.1%支付滞纳金。')

    doc.add_heading('第三条 保密条款', level=1)
    doc.add_paragraph('3.1 双方应对在合同履行过程中知悉的对方商业秘密、技术秘密及其他未公开信息承担保密义务。')
    doc.add_paragraph('3.2 未经对方书面同意，任何一方不得向第三方泄露。')
    doc.add_paragraph('3.3 本保密义务在合同终止后三年内仍然有效。')

    doc.add_heading('第四条 违约责任', level=1)
    doc.add_paragraph('4.1 当事人一方不履行合同义务或者履行合同义务不符合约定的，应当承担继续履行、采取补救措施或者赔偿损失等违约责任。')
    doc.add_paragraph('4.2 若甲方逾期支付款项超过30日的，乙方有权终止服务。')

    doc.save(output_path)
    print(f"已创建: {output_path}")


def create_contract_3(output_path):
    doc = Document()
    doc.add_heading('劳动合同', 0)

    doc.add_heading('第一条 试用期', level=1)
    doc.add_paragraph('1.1 本合同试用期为三个月，自入职之日起计算。')
    doc.add_paragraph('1.2 试用期内，如乙方不符合录用条件，甲方有权解除本合同。')

    doc.add_heading('第二条 工作时间', level=1)
    doc.add_paragraph('2.1 乙方实行标准工时制，每日工作8小时，每周工作40小时。')
    doc.add_paragraph('2.2 乙方享有国家规定的法定节假日。')

    doc.add_heading('第三条 劳动报酬', level=1)
    doc.add_paragraph('3.1 乙方月工资为人民币15000元整。')
    doc.add_paragraph('3.2 甲方应于每月10日前支付上月工资。')

    doc.add_heading('第四条 保密和竞业限制', level=1)
    doc.add_paragraph('4.1 乙方应对在工作中知悉的甲方商业秘密承担保密义务。')
    doc.add_paragraph('4.2 劳动合同终止或解除后两年内，乙方不得在与甲方有竞争关系的单位任职。')

    doc.add_heading('第五条 合同解除', level=1)
    doc.add_paragraph('5.1 经双方协商一致，可以解除本合同。')
    doc.add_paragraph('5.2 乙方提前30日以书面形式通知甲方，可以解除劳动合同。')

    doc.save(output_path)
    print(f"已创建: {output_path}")


if __name__ == '__main__':
    test_dir = os.path.join(os.path.dirname(__file__), 'sample_data', 'contracts')
    os.makedirs(test_dir, exist_ok=True)

    create_contract_1(os.path.join(test_dir, '房屋买卖合同.docx'))
    create_contract_2(os.path.join(test_dir, '服务合同.docx'))
    create_contract_3(os.path.join(test_dir, '劳动合同.docx'))

    print(f"\n测试文档已创建在: {test_dir}")
