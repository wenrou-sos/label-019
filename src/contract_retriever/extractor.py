import os
import re
from pathlib import Path
from dataclasses import dataclass, field
from typing import List, Optional, Tuple, Any
from tqdm import tqdm
import logging

logger = logging.getLogger(__name__)


@dataclass
class TextSegment:
    text: str
    page: int
    paragraph: int
    start_pos: int = 0
    end_pos: int = 0


@dataclass
class DocumentContent:
    file_path: str
    file_type: str
    segments: List[TextSegment] = field(default_factory=list)
    metadata: dict = field(default_factory=dict)
    error: Optional[str] = None

    @property
    def full_text(self) -> str:
        return "\n".join([seg.text for seg in self.segments])

    @property
    def total_pages(self) -> int:
        if not self.segments:
            return 0
        return max(seg.page for seg in self.segments)

    @property
    def total_paragraphs(self) -> int:
        return len(self.segments)


class BaseExtractor:
    def extract(self, file_path: str) -> DocumentContent:
        raise NotImplementedError


class PDFExtractor(BaseExtractor):
    def __init__(self):
        try:
            import pdfplumber
            self.pdfplumber = pdfplumber
        except ImportError:
            logger.error("pdfplumber not installed. Please install it with: pip install pdfplumber")
            raise

    def extract(self, file_path: str) -> DocumentContent:
        content = DocumentContent(file_path=file_path, file_type="pdf")
        try:
            with self.pdfplumber.open(file_path) as pdf:
                content.metadata = {
                    "num_pages": len(pdf.pages),
                    "metadata": pdf.metadata,
                }
                paragraph_count = 0
                global_pos = 0

                for page_num, page in enumerate(pdf.pages, start=1):
                    page_text = page.extract_text() or ""
                    tables = page.extract_tables() or []

                    if page_text.strip():
                        paragraphs = self._split_paragraphs(page_text)
                        for para_text in paragraphs:
                            if para_text.strip():
                                segment = TextSegment(
                                    text=para_text.strip(),
                                    page=page_num,
                                    paragraph=paragraph_count,
                                    start_pos=global_pos,
                                    end_pos=global_pos + len(para_text)
                                )
                                content.segments.append(segment)
                                paragraph_count += 1
                                global_pos += len(para_text) + 1

                    for table_idx, table in enumerate(tables):
                        table_text = self._table_to_text(table)
                        if table_text.strip():
                            segment = TextSegment(
                                text=f"[表格 {page_num}-{table_idx + 1}]\n{table_text}",
                                page=page_num,
                                paragraph=paragraph_count,
                                start_pos=global_pos,
                                end_pos=global_pos + len(table_text)
                            )
                            content.segments.append(segment)
                            paragraph_count += 1
                            global_pos += len(table_text) + 1

        except Exception as e:
            content.error = f"PDF解析失败: {str(e)}"
            logger.error(f"Error extracting PDF {file_path}: {e}")

        return content

    def _split_paragraphs(self, text: str) -> List[str]:
        paragraphs = re.split(r'\n\s*\n', text)
        return [p.strip() for p in paragraphs if p.strip()]

    def _table_to_text(self, table: List[List[Any]]) -> str:
        rows = []
        for row in table:
            cells = [str(cell).strip() if cell is not None else "" for cell in row]
            rows.append(" | ".join(cells))
        return "\n".join(rows)


class WordExtractor(BaseExtractor):
    def __init__(self):
        try:
            from docx import Document
            self.Document = Document
        except ImportError:
            logger.error("python-docx not installed. Please install it with: pip install python-docx")
            raise

    def extract(self, file_path: str) -> DocumentContent:
        content = DocumentContent(file_path=file_path, file_type="docx")
        try:
            doc = self.Document(file_path)
            content.metadata = {
                "num_paragraphs": len(doc.paragraphs),
                "num_tables": len(doc.tables),
                "core_properties": {
                    "title": doc.core_properties.title,
                    "author": doc.core_properties.author,
                    "created": str(doc.core_properties.created) if doc.core_properties.created else None,
                    "modified": str(doc.core_properties.modified) if doc.core_properties.modified else None,
                }
            }

            paragraph_count = 0
            page_number = 1
            global_pos = 0

            for para in doc.paragraphs:
                if para.text.strip():
                    if self._is_page_break(para):
                        page_number += 1
                        continue

                    segment = TextSegment(
                        text=para.text.strip(),
                        page=page_number,
                        paragraph=paragraph_count,
                        start_pos=global_pos,
                        end_pos=global_pos + len(para.text)
                    )
                    content.segments.append(segment)
                    paragraph_count += 1
                    global_pos += len(para.text) + 1

            for table_idx, table in enumerate(doc.tables):
                table_text = self._table_to_text(table)
                if table_text.strip():
                    segment = TextSegment(
                        text=f"[表格 {table_idx + 1}]\n{table_text}",
                        page=page_number,
                        paragraph=paragraph_count,
                        start_pos=global_pos,
                        end_pos=global_pos + len(table_text)
                    )
                    content.segments.append(segment)
                    paragraph_count += 1
                    global_pos += len(table_text) + 1

        except Exception as e:
            content.error = f"Word文档解析失败: {str(e)}"
            logger.error(f"Error extracting Word document {file_path}: {e}")

        return content

    def _is_page_break(self, para) -> bool:
        for run in para.runs:
            if 'w:lastRenderedPageBreak' in run._element.xml:
                return True
        return False

    def _table_to_text(self, table) -> str:
        rows = []
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            rows.append(" | ".join(cells))
        return "\n".join(rows)


class DocumentScanner:
    SUPPORTED_EXTENSIONS = {'.pdf', '.docx', '.doc'}

    def __init__(self, directory: str):
        self.directory = directory
        self.extractors = {
            '.pdf': PDFExtractor(),
            '.docx': WordExtractor(),
            '.doc': WordExtractor(),
        }

    def scan(self, show_progress: bool = True) -> Tuple[List[DocumentContent], dict]:
        files = self._find_files()
        results = []
        stats = {
            'total_files': len(files),
            'pdf_count': 0,
            'docx_count': 0,
            'doc_count': 0,
            'success_count': 0,
            'failed_files': [],
            'extraction_time_ms': 0,
        }

        import time
        start_time = time.time()

        iterator = tqdm(files, desc="提取文档内容", disable=not show_progress) if show_progress else files

        for file_path in iterator:
            ext = Path(file_path).suffix.lower()
            if ext == '.pdf':
                stats['pdf_count'] += 1
            elif ext == '.docx':
                stats['docx_count'] += 1
            elif ext == '.doc':
                stats['doc_count'] += 1

            extractor = self.extractors.get(ext)
            if extractor:
                content = extractor.extract(file_path)
                if content.error:
                    stats['failed_files'].append(file_path)
                else:
                    stats['success_count'] += 1
                results.append(content)

        stats['extraction_time_ms'] = int((time.time() - start_time) * 1000)

        return results, stats

    def _find_files(self) -> List[str]:
        files = []
        for root, _, filenames in os.walk(self.directory):
            for filename in filenames:
                ext = Path(filename).suffix.lower()
                if ext in self.SUPPORTED_EXTENSIONS:
                    files.append(os.path.join(root, filename))
        return sorted(files)


def extract_contract_clauses(content: DocumentContent) -> List[dict]:
    clauses = []
    clause_patterns = [
        r'^第[一二三四五六七八九十百千\d]+[章节条款条]',
        r'^[一二三四五六七八九十百千]+[、\.]\s+',
        r'^\d+[、\.]\s+',
        r'^(定义|释义|总则|目的|范围)\s*[:：]?\s*$',
    ]

    current_clause_title = None
    current_clause_content = []
    current_page = 1
    current_paragraph = 0

    for seg in content.segments:
        text = seg.text.strip()
        if not text:
            continue

        is_clause_start = False
        for pattern in clause_patterns:
            if re.search(pattern, text):
                is_clause_start = True
                break

        if is_clause_start and len(text) < 80:
            if current_clause_title and current_clause_content:
                clauses.append({
                    'title': current_clause_title,
                    'content': '\n'.join(current_clause_content),
                    'page': current_page,
                    'paragraph': current_paragraph,
                    'file_path': content.file_path,
                })
            current_clause_title = text
            current_clause_content = []
            current_page = seg.page
            current_paragraph = seg.paragraph
        else:
            if current_clause_title:
                current_clause_content.append(text)
            current_page = seg.page

    if current_clause_title and current_clause_content:
        clauses.append({
            'title': current_clause_title,
            'content': '\n'.join(current_clause_content),
            'page': current_page,
            'paragraph': current_paragraph,
            'file_path': content.file_path,
        })

    return clauses
